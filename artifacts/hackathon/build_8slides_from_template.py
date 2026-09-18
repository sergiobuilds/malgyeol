#!/usr/bin/env python3
"""말결 정책·기술 발표 8장 — 비보북 양식(말결 (+3)의 사본.pptx) 기반 제작."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from lxml import etree

TPL = sys.argv[1]
OUT = sys.argv[2]

IVORY = RGBColor(0xF4, 0xF1, 0xEB)
GRAY = RGBColor(0xC4, 0xC3, 0xBF)
ORANGE = RGBColor(0xF4, 0xA0, 0x4B)
RULE = RGBColor(0x45, 0x46, 0x47)
PANEL = RGBColor(0x24, 0x26, 0x29)
FONT = "맑은 고딕"

prs = Presentation(TPL)

# ---- 1. 슬라이드 9·10 제거 (빈 Slide #8, DEMO) -------------------------------
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[8:]:
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
assert len(prs.slides) == 8

# ---- 2. 양식 footer(하단 행)만 남기고 본문 도형 제거, 노트 비우기 -------------
for s in prs.slides:
    for sh in list(s.shapes):
        if sh.top < Inches(5.0):
            sh._element.getparent().remove(sh._element)
    if s.has_notes_slide:
        s.notes_slide.notes_text_frame.text = ""


# ---- helpers -------------------------------------------------------------------
def _set_font(run, size, color, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)
    rPr.set("lang", "ko-KR")


def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       margin=0.04, line_spacing=1.12, space_after=2):
    """paras: list of (text, size, color, bold) or list of list-of-runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for c in list(bodyPr):
        bodyPr.remove(c)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = align
        para.line_spacing = line_spacing
        para.space_after = Pt(space_after)
        runs = p if isinstance(p, list) else [p]
        for (text, size, color, bold) in runs:
            r = para.add_run()
            r.text = text
            _set_font(r, size, color, bold)
    return box


def panel(slide, x, y, w, h, fill=PANEL, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def line(slide, x1, y1, x2, y2, color=RULE, width=0.75, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def title(slide, text, sub):
    tb(slide, 0.55, 0.58, 7.8, 0.5, [(text, 24, IVORY, True)])
    tb(slide, 0.55, 1.08, 8.4, 0.34, [(sub, 13, GRAY, False)])


def label(slide, x, y, w, text, color=ORANGE, size=11):
    tb(slide, x, y, w, 0.28, [(text, size, color, True)])


def body(slide, x, y, w, h, lines, size=13, color=IVORY, sa=3, ls=1.12):
    tb(slide, x, y, w, h, [(t, size, color, False) for t in lines], space_after=sa, line_spacing=ls)


def footer(slide, text, y=4.42):
    tb(slide, 0.55, y, 8.9, 0.32, [[("▎", 12, ORANGE, True), (text, 12, IVORY, False)]])


def source(slide, text):
    tb(slide, 0.55, 4.78, 8.9, 0.26, [(text, 8, GRAY, False)])


S = list(prs.slides)

# =============================== 1 표지 + 스토리텔링 ===============================
s = S[0]
tb(s, 0.5, 0.72, 4.4, 0.85, [("말결", 40, IVORY, True)])
tb(s, 0.55, 1.55, 4.4, 0.4, [("전화 기반 생활지원 연계", 16, IVORY, False)])
label(s, 0.55, 2.2, 3, "핵심 사례")
body(s, 0.55, 2.47, 4.4, 0.7, ["쌀 보유·조리 곤란", "생필품 부족·방문 곤란"], size=14)
body(s, 0.55, 3.2, 4.4, 0.7, ["복합 필요와 분리된 지원 경로", "기관 탐색·반복 설명·조건 조율"], size=12, color=GRAY)
# 우측: 관계 구조(지도 그래픽 위 면 구성)
px, py, pw, ph = 5.35, 0.95, 4.05, 3.15
panel(s, px, py, pw, ph, line=RULE)
rows = [("생활 언어", "시민의 필요·이동·조리 제약"),
        ("기관 연결", "사업별 이용 조건·접수 창구"),
        ("후속 처리", "기관 문의·시민 회신·재선택")]
for i, (h, d) in enumerate(rows):
    ry = py + 0.18 + i * 1.0
    tb(s, px + 0.2, ry, 1.4, 0.32, [(f"{i+1}  {h}", 13, ORANGE, True)])
    tb(s, px + 0.2, ry + 0.33, pw - 0.4, 0.34, [(d, 12.5, IVORY, False)])
    if i < 2:
        line(s, px + 0.42, ry + 0.72, px + 0.42, ry + 0.95, arrow=True)
panel(s, 0.45, 4.22, 5.0, 0.42, fill=RGBColor(0, 0, 0))
tb(s, 0.55, 4.27, 4.9, 0.32, [[("▎", 11.5, ORANGE, True), ("전화 한 통의 접수 · 복수 필요의 분리 · 지원 경로의 연결", 11.5, IVORY, False)]])
tb(s, 0.55, 4.78, 4.2, 0.26, [("출처: POLICY_BASIS.md 1·2절, 2026.9.18 확인", 8, GRAY, False)])

# =============================== 2 법령 근거와 구체화 ===============================
s = S[1]
title(s, "정책 근거의 업무 구체화", "통합지원의 제도 기반 · 상담 이후의 실행 부담")
cy, ch = 1.55, 2.75
# 열 좌표
c1x, c1w = 0.55, 2.15
c2x, c2w = 2.95, 1.95
c3x, c3w = 5.15, 1.95
c4x, c4w = 7.35, 2.1
label(s, c1x, cy, c1w, "근거 영역")
label(s, c2x, cy, c2w, "제도적 업무")
label(s, c3x, cy, c3w, "제품의 실행 단위")
label(s, c4x, cy, c4w, "권한 구분")
line(s, c1x, cy + 0.32, 9.45, cy + 0.32)
body(s, c1x, cy + 0.42, c1w, 1.6, ["돌봄통합지원법", "돌봄통합지원법 시행규칙", "서울시 지역사회 돌봄\n통합지원 조례"], size=13, sa=6)
rowsL = ["지역 자원·전달체계", "서비스 연계·점검", "신청·조사·개인별지원계획"]
rowsR = ["사업별 접수처·이용 절차", "기관 문의·일정·수령 조건", "신청 의사 전달·결과 회신"]
for i in range(3):
    ry = cy + 0.42 + i * 0.62
    tb(s, c2x, ry, c2w, 0.5, [(rowsL[i], 13, IVORY, False)], anchor=MSO_ANCHOR.MIDDLE)
    line(s, c2x + c2w + 0.03, ry + 0.25, c3x - 0.05, ry + 0.25, color=ORANGE, arrow=True)
    tb(s, c3x, ry, c3w, 0.5, [(rowsR[i], 13, IVORY, False)], anchor=MSO_ANCHOR.MIDDLE)
line(s, c4x - 0.12, cy + 0.42, c4x - 0.12, cy + ch - 0.2)
body(s, c4x, cy + 0.42, c4w, 1.9, ["시민의 동의·선택", "기관의 자격·제공 판단", "담당자의 예외 처리·인계"], size=13, sa=6)
footer(s, "법령 근거 → 허용 업무 → 동의 범위 → 제한 실증")
source(s, "출처: 의료ㆍ요양 등 지역 돌봄의 통합지원에 관한 법률(법률 제21777호, 2026.9.10 시행) · 동법 시행규칙(2026.3.27 시행) · "
          "서울특별시 지역사회 돌봄 통합지원에 관한 조례(2026.1.5 시행), 국가법령정보센터 2026.9.18 확인")

# =============================== 3 팀 소개와 문제의식 ===============================
s = S[2]
title(s, "디지털 약자의 AI 접근성", "팀의 전문영역 · 익숙한 전화 기반 접점")
lx, lw = 0.55, 5.0
rx, rw = 6.35, 3.1
ay = 1.55
label(s, lx, ay, lw, "전문영역")
label(s, rx, ay, rw, "이용 접점")
line(s, lx, ay + 0.32, 9.45, ay + 0.32)
blocks = [("A  정책·회계", ["사업 절차·규정 해석", "업무 권한·증거·책임", "현장 실증·성과 측정"]),
          ("B  AI·에이전트", ["요구 명확화·작업 분해", "역할별 도구·검증 절차", "동의·상태·이력 관리"])]
for i, (h, items) in enumerate(blocks):
    by = ay + 0.45 + i * 1.22
    panel(s, lx, by, lw, 1.1)
    tb(s, lx + 0.15, by + 0.08, 1.6, 0.32, [(h, 13, ORANGE, True)])
    body(s, lx + 1.85, by + 0.08, lw - 2.0, 1.0, items, size=12.5, sa=2)
    line(s, lx + lw, by + 0.55, rx - 0.08, ay + 0.45 + 1.15, color=RULE, arrow=(i == 1))
panel(s, rx, ay + 0.45, rw, 2.32, line=RULE)
body(s, rx + 0.15, ay + 0.6, rw - 0.3, 2.1,
     ["시민의 음성 요청", "기관의 전화 응답", "담당자의 미해결 요청 관리"], size=13, sa=10)
footer(s, "앱 설치·검색·반복 입력의 부담 완화 · 시민과 기관의 기존 접점")
source(s, "출처: 말결 정책·실증 근거(POLICY_BASIS.md) 2·9절 · 계획 인계문(planning-handoff.md), 2026.9.18 확인")

# =============================== 4 정책 4개 연결 ===============================
s = S[3]
title(s, "4개 사업의 생활지원 연결", "식사·생필품 중심 범위 · 사업별 기준과 접수 경로")
progs = [("푸드뱅크·푸드마켓", "식품·생필품 지원", "선정·수령 조건"),
         ("찾아가는 푸드마켓", "지역별 운영", "이동 제약·수령 방식"),
         ("그냥드림", "위기 식품지원", "방문·상담 절차"),
         ("돌봄SOS", "식사배달·돌봄", "상담·기관 의뢰")]
gx, gw, gap, gy, gh = 0.55, 2.14, 0.113, 1.55, 1.62
for i, (n, scope, chk) in enumerate(progs):
    x = gx + i * (gw + gap)
    panel(s, x, gy, gw, gh)
    tb(s, x + 0.12, gy + 0.08, gw - 0.2, 0.3, [[(f"{i+1} ", 12, ORANGE, True), (n, 12.5, IVORY, True)]])
    tb(s, x + 0.12, gy + 0.48, gw - 0.2, 0.26, [("지원 범위", 10, GRAY, False)])
    tb(s, x + 0.12, gy + 0.7, gw - 0.2, 0.3, [(scope, 12.5, IVORY, False)])
    tb(s, x + 0.12, gy + 1.02, gw - 0.2, 0.26, [("확인 항목", 10, GRAY, False)])
    tb(s, x + 0.12, gy + 1.24, gw - 0.2, 0.3, [(chk, 12.5, IVORY, False)])
by = 3.32
label(s, 0.55, by, 3, "공통 연결 구조")
tb(s, 0.55, by + 0.28, 5.2, 0.34, [("사업 ↔ 기관·사업장 ↔ 관할·목적별 창구 ↔ 이용 절차", 12.5, IVORY, False)])
line(s, 6.05, by + 0.05, 6.05, by + 0.62)
label(s, 6.25, by, 3.2, "확인된 데이터 규모")
tb(s, 6.25, by + 0.28, 3.2, 0.34, [[("공식 원천 65행 → 정규화 37기관 + 보완 3창구 = ", 11.5, IVORY, False), ("40개 기관·창구", 11.5, ORANGE, True)]])
footer(s, "필요별 경로 조회 · 기관별 조건 확인 · 후속 문의", y=4.1)
source(s, "출처: 서울잇다푸드뱅크 이용안내·기관 목록·그냥드림 안내 · 서울시 찾아가는 푸드마켓 운영(2026.3.20 수정) · 서울복지포털 돌봄SOS · "
          "말결 data/support-network.json, 2026.9.18 확인")

# =============================== 5 프록시 방식 ===============================
s = S[4]
title(s, "동의 기반 대리 수행", "시민 요청·기관 문의·결과 회신의 실행 구조")
stages = [("시민 접수", ["필요·제약 분리", "기관·목적·전달 정보 동의"]),
          ("기관 문의", ["이용 조건·일정 확인", "동의 범위의 신청 의사 전달"]),
          ("시민 회신", ["가능 조건·준비사항", "중요 변경의 재선택"]),
          ("후속 처리", ["기관 재연락", "미해결 필요·담당자 인계"])]
sx, sw, sgap, sy, shh = 0.55, 2.0, 0.3, 1.6, 1.45
for i, (n, items) in enumerate(stages):
    x = sx + i * (sw + sgap)
    panel(s, x, sy, sw, shh)
    tb(s, x + 0.12, sy + 0.08, sw - 0.2, 0.32, [[(f"{i+1}  ", 12, ORANGE, True), (n, 13, IVORY, True)]])
    body(s, x + 0.12, sy + 0.45, sw - 0.2, 0.95, items, size=12, sa=3)
    if i < 3:
        line(s, x + sw + 0.04, sy + 0.3, x + sw + sgap - 0.04, sy + 0.3, color=ORANGE, arrow=True)
# 동의 조건 표시 (1→2 화살표 아래)
tb(s, sx + sw - 0.05, sy + 0.4, sgap + 0.1, 0.24, [("동의", 8.5, ORANGE, False)], align=PP_ALIGN.CENTER)
# 재선택 → 기관 재문의 피드백 (3 → 2)
b2c = sx + 1 * (sw + sgap) + sw / 2
b3c = sx + 2 * (sw + sgap) + sw / 2
fy = sy + shh + 0.22
line(s, b3c, sy + shh, b3c, fy)
line(s, b3c, fy, b2c, fy)
line(s, b2c, fy, b2c, sy + shh + 0.02, arrow=True)
tb(s, b2c + 0.15, fy + 0.01, b3c - b2c - 0.3, 0.24, [("시민 재선택 → 기관 재문의", 9.5, GRAY, False)], align=PP_ALIGN.CENTER)
label(s, 0.55, 3.55, 2.2, "부분 해결 사례")
tb(s, 2.0, 3.55, 6.5, 0.3, [("식사 연결 유지 · 생필품 경로 재탐색", 12.5, IVORY, False)])
footer(s, "문의·신청 의사 전달·접수·제공의 구분 · 기관의 최종 판단", y=4.05)
source(s, "출처: 말결 정책·실증 근거(POLICY_BASIS.md) 5·7.2절 · 전화 처리 계약(phone-goal-prompt.md), 2026.9.18 확인")

# =============================== 6 우로보로스 철학과 기술 ===============================
s = S[5]
title(s, "질문·구조화·검증의 설계 원리", "생활 발화의 명확화 · 완료 조건 중심의 작업 분해")
cols = [("질문", "요구의 명확화", ["필요·제약·선호의 구분", "모호한 발화의 추가 확인", "시민 선택·완료 조건"]),
        ("구조화", "업무의 분해", ["요청 아래 복수 필요", "필요별 기관·문의·통화", "역할별 도구·허용 범위"]),
        ("검증", "결과의 재확인", ["기관 답변·조건의 기록", "선택 변경·후속 문의", "부분 해결·미해결 관리"])]
ox, ow, ogap, oy, oh = 0.55, 2.7, 0.4, 1.6, 1.85
for i, (h, sub, items) in enumerate(cols):
    x = ox + i * (ow + ogap)
    panel(s, x, oy, ow, oh)
    tb(s, x + 0.15, oy + 0.1, ow - 0.3, 0.34, [[(h, 15, ORANGE, True), ("   " + sub, 12, GRAY, False)]])
    line(s, x + 0.15, oy + 0.5, x + ow - 0.15, oy + 0.5)
    body(s, x + 0.15, oy + 0.58, ow - 0.3, 1.2, items, size=12.5, sa=4)
    if i < 2:
        line(s, x + ow + 0.05, oy + 0.27, x + ow + ogap - 0.05, oy + 0.27, color=ORANGE, arrow=True)
# 검증 → 후속 질문 순환
c1 = ox + ow / 2
c3 = ox + 2 * (ow + ogap) + ow / 2
ly = oy + oh + 0.25
line(s, c3, oy + oh, c3, ly)
line(s, c3, ly, c1, ly)
line(s, c1, ly, c1, oy + oh + 0.02, arrow=True)
tb(s, c1 + 0.3, ly + 0.01, c3 - c1 - 0.6, 0.26, [("실행 결과에 따른 후속 질문", 9.5, GRAY, False)], align=PP_ALIGN.CENTER)
footer(s, "질문 → 구조화 → 실행 → 검증 → 후속 질문", y=4.1)
source(s, "출처: 말결 설계 원리의 적용 · scripts/coordination_tools.py 역할별 도구·동의·선택 기록, 2026.9.18 기준")

# =============================== 7 회계 DSL 기술 ===============================
s = S[6]
title(s, "규범 구조화와 실행 통제", "회계 DSL 기반기술 · 말결의 동의·상태 검사")
lx, lw, rx, rw, py, ph = 0.55, 4.2, 5.25, 4.2, 1.55, 2.7
panel(s, lx, py, lw, ph)
panel(s, rx, py, rw, ph)
tb(s, lx + 0.15, py + 0.08, lw - 0.3, 0.3, [[("회계 DSL의 규범 컴파일", 13, ORANGE, True), ("   기반기술", 10.5, GRAY, False)]])
line(s, lx + 0.15, py + 0.44, lx + lw - 0.15, py + 0.44)
chain = ["원문 조항", "범위·효력 시점", "주체·행위·조건", "의무·허용·금지", "정적 검사·규범 충돌", "서명 규정팩"]
for i, t in enumerate(chain):
    ry = py + 0.52 + i * 0.27
    tb(s, lx + 0.15, ry, 0.35, 0.26, [(f"{i+1}", 11, ORANGE, True)])
    tb(s, lx + 0.5, ry, lw - 0.7, 0.26, [(t, 12, IVORY, False)])
line(s, lx + 0.3, py + 0.78, lx + 0.3, py + 0.52 + 5 * 0.27 - 0.02)
tb(s, lx + 0.15, py + 2.12, lw - 0.3, 0.24, [("보조 항목", 10, GRAY, False)])
tb(s, lx + 0.15, py + 2.34, lw - 0.3, 0.3, [("판정 근거 보존 · 규정 변경 영향 · 회귀 검증", 11.5, IVORY, False)])
tb(s, rx + 0.15, py + 0.08, rw - 0.3, 0.3, [[("말결의 실행 시점 검사", 13, ORANGE, True), ("   현행 코드", 10.5, GRAY, False)]])
line(s, rx + 0.15, py + 0.44, rx + rw - 0.15, py + 0.44)
checks = ["기관·목적·전달 필드의 동의 범위", "요청 revision·중단 상태", "멱등키·중복 실행 방지", "결과 불명 시 자동 재시도 차단",
          "요청·문의·시도·답변의 연결", "시민 선택·후속 행동 기록"]
for i, t in enumerate(checks):
    ry = py + 0.52 + i * 0.35
    tb(s, rx + 0.15, ry, 0.35, 0.3, [("▪", 10, ORANGE, False)])
    tb(s, rx + 0.45, ry, rw - 0.65, 0.3, [(t, 12, IVORY, False)])
line(s, lx + lw + 0.08, py + ph / 2, rx - 0.08, py + ph / 2, color=ORANGE, arrow=True)
tb(s, lx + lw, py + ph / 2 - 0.36, rx - lx - lw, 0.26, [("적용", 9.5, ORANGE, False)], align=PP_ALIGN.CENTER)
footer(s, "규칙·근거·실행 이력의 연결 · 적용 범위별 검증")
source(s, "출처: 회계 DSL 기술도식(규정 위계·규범 컴파일·선택적 판정 게이트·회귀 검증) · src/coordination/engine.ts 검사 조건, 2026.9.18 기준")

# =============================== 8 상세 기술 파이프라인 ===============================
s = S[7]
title(s, "음성·업무·기록의 연결", "역할별 세션과 공통 API · 상태 검사와 영속 저장")
steps = [("전화·음성", ["ClawOps", "Gemini 실시간 음성"]),
         ("역할별 도구", ["시민 접수", "기관 문의", "시민 회신"]),
         ("공통 API", ["지원망 조회", "동의·문의·선택"]),
         ("업무 엔진", ["요청 버전·권한", "필요·문의 상태"]),
         ("영속 기록", ["SQLite 요청 원장", "음성 작업 journal"])]
px0, pw, pgap, py0, ph0 = 0.55, 1.68, 0.125, 1.5, 1.22
for i, (n, items) in enumerate(steps):
    x = px0 + i * (pw + pgap)
    panel(s, x, py0, pw, ph0)
    tb(s, x + 0.1, py0 + 0.07, pw - 0.2, 0.3, [[(f"{i+1}  ", 11, ORANGE, True), (n, 12, IVORY, True)]])
    body(s, x + 0.08, py0 + 0.42, pw - 0.14, 0.78, items, size=10.5, sa=2)
    if i < 4:
        line(s, x + pw + 0.01, py0 + 0.22, x + pw + pgap - 0.01, py0 + 0.22, color=ORANGE, arrow=True)
ctrl_y = 2.9
label(s, 0.55, ctrl_y, 4, "실행 통제")
line(s, 0.55, ctrl_y + 0.3, 9.45, ctrl_y + 0.3)
ctrls = [("발신 전", ["기관 동의·허용 전달 필드", "중단·오래된 revision 차단"]),
         ("통화 중·종료 후", ["멱등키·활성 통화 경쟁 통제", "종료 결과·기관 답변 대조"]),
         ("시민 회신", ["수신자 확인·중요 조건 선택", "후속 문의·담당자 인계"])]
cw, cg = 2.85, 0.175
for i, (h, items) in enumerate(ctrls):
    x = 0.55 + i * (cw + cg)
    tb(s, x, ctrl_y + 0.38, cw, 0.28, [(h, 12, IVORY, True)])
    body(s, x, ctrl_y + 0.66, cw, 0.62, items, size=11.5, sa=2)
    if i > 0:
        line(s, x - cg / 2, ctrl_y + 0.42, x - cg / 2, ctrl_y + 1.25)
footer(s, "BEGIN IMMEDIATE · WAL / FULL · 결과 불명 시 자동 재시도 차단", y=4.33)
source(s, "출처: src/coordination/engine.ts·store.ts · scripts/coordination_voice.py·coordination_tools.py · docs/ARCHITECTURE.md, 2026.9.18 기준")

prs.save(OUT)
print("saved", OUT, len(prs.slides))
